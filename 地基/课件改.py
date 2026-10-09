
import json
import uuid
import os

from 地基 import 大模型, 提示词, 课件

能改的字段 = ("标题", "关键句", "要点", "备注")
要点前缀 = "·  "
摘要上限 = 80


def 读结构(课件路径):
    结构 = []
    for 一页 in 课件.读各页(课件路径):
        备注 = 一页["备注"]
        结构.append({"页": 一页["页"], "标题": 一页["标题"], "关键句": 一页["关键句"],
                   "要点": 一页["要点"],
                   "备注": (备注[:摘要上限] + "…") if len(备注) > 摘要上限 else 备注})
    return 结构


def 结构文字(结构):
    行 = []
    for 一页 in 结构:
        行.append("第 %d 页：标题=%s ｜ 关键句=%s" % (一页["页"], 一页["标题"] or "（空）",
                                              一页["关键句"] or "（空）"))
        行.append("    要点：%s" % (一页["要点"] or "（空）").replace("\n", " "))
        行.append("    讲稿（备注）：%s" % (一页["备注"] or "（空）"))
    return "\n".join(行)


def 出清单(课件路径, 老师的话, 服务商=None):
    结构 = 读结构(课件路径)
    原文 = 提示词.填充(
        提示词.读取("改课件"),
        课件结构=结构文字(结构),
        老师的话=str(老师的话 or "").strip(),
    )
    回复 = 大模型.问(原文, 服务商=服务商)
    return 解析清单(回复, 页数=len(结构))


def 解析清单(原文, 页数):
    try:
        数据 = 课件.抠出JSON对象(原文, "改课件的回复")
    except Exception as 异常:
        raise ValueError("模型这次没给出能用的改动清单（回复不是 JSON）："
                         + type(异常).__name__ + "：" + str(异常)[:60]) from 异常
    if not isinstance(数据, dict):
        raise ValueError("模型给的清单不是一份对象（{}），没法当改动用")
    改动们 = 数据.get("改动")
    if not isinstance(改动们, list) or not 改动们:
        return {"改动": [], "做不到的": [str(x) for x in (数据.get("做不到的") or [])]}
    改动, 做不到 = [], list(str(x) for x in (数据.get("做不到的") or []))
    for 一条 in 改动们:
        if not isinstance(一条, dict):
            做不到.append("有一项改动不是对象，已跳过：" + str(一条)[:40])
            continue
        try:
            号 = int(一条.get("页"))
        except (TypeError, ValueError):
            做不到.append("有一项改动没写清是第几页，已跳过：" + json.dumps(一条, ensure_ascii=False)[:40])
            continue
        if 号 < 1 or 号 > 页数:
            做不到.append("第 %d 页不存在（这份课件一共 %d 页），已跳过" % (号, 页数))
            continue
        这页 = {}
        for 字段 in 能改的字段:
            if 字段 not in 一条:
                continue
            值 = 一条[字段]
            if 字段 == "要点":
                if isinstance(值, str):
                    值 = [x.strip() for x in 值.split("\n") if x.strip()]
                if not isinstance(值, list) or not all(isinstance(x, str) for x in 值):
                    做不到.append("第 %d 页的要点格式不对（要一串短句），已跳过" % 号)
                    continue
                干净 = [x.strip() for x in 值 if x.strip()]
                if not 干净:
                    做不到.append("第 %d 页的要点是空的，已跳过（要清空要点请在 WPS 里做）" % 号)
                    continue
                这页[字段] = 干净
            else:
                if not isinstance(值, str):
                    值 = str(值)
                这页[字段] = 值.strip()
        多出来的 = [k for k in 一条 if k != "页" and k not in 能改的字段]
        if 多出来的:
            做不到.append("第 %d 页里有这一版改不了的项（%s）—— "
                        "这一版只能改标题/关键句/要点/讲稿，其余请在 WPS 里改"
                        % (号, "、".join(多出来的)))
        if 这页:
            这页["页"] = 号
            改动.append(这页)
    return {"改动": 改动, "做不到的": 做不到}


def 预览清单(清单):
    if not 清单["改动"]:
        行 = ["这次**没有要改的地方**。"]
    else:
        行 = ["我打算改这几处："]
        for 一条 in 清单["改动"]:
            行.append("· 第 %d 页：" % 一条["页"])
            for 字段 in 能改的字段:
                if 字段 not in 一条:
                    continue
                if 字段 == "要点":
                    行.append("    要点改成：" + " / ".join(x.replace(chr(10), " ") for x in 一条[字段]))
                else:
                    行.append("    %s 改成：%s" % (字段, 一条[字段].replace(chr(10), " ")[:60]))
    if 清单["做不到的"]:
        行.append("这些我做不到（请自己在 WPS 里处理）：")
        行.extend("· " + x for x in 清单["做不到的"])
    return "\n".join(行)


def 要改的页(清单):
    return sorted({一条["页"] for 一条 in 清单["改动"]})


def 执行清单(课件路径, 清单, 警告=None):
    from pptx import Presentation
    好页, 没成 = [], []
    if not 清单.get("改动"):
        return 好页, 没成

    备份名 = 课件.备份成(课件路径, "改前")
    演示文稿 = Presentation(str(课件路径))
    页们 = list(演示文稿.slides)

    def 按名字找(页, 名):
        for 形 in 页.shapes:
            try:
                if 形.name == 名 and 形.has_text_frame:
                    return 形
            except Exception:
                continue
        return None

    def 抄格式(框):
        档 = {"名字": None, "字号": None, "颜色": None}
        try:
            第一 = 框.paragraphs[0].runs[0]
            档 = {"名字": 第一.font.name, "字号": 第一.font.size,
                  "颜色": 第一.font.color.rgb if 第一.font.color and 第一.font.color.type else None}
        except Exception:
            pass
        return 档

    def 写回(框, 段落们):
        档 = 抄格式(框)
        框.clear()
        for 序号, 一段 in enumerate(段落们):
            段 = 框.paragraphs[0] if 序号 == 0 else 框.add_paragraph()
            段.text = str(一段)
            for 片段 in 段.runs:
                if 档["名字"]:
                    片段.font.name = 档["名字"]
                if 档["字号"] is not None:
                    片段.font.size = 档["字号"]
                if 档["颜色"] is not None:
                    片段.font.color.rgb = 档["颜色"]

    for 一条 in 清单["改动"]:
        页 = 页们[一条["页"] - 1]
        这页成了 = False
        for 字段, 值 in 一条.items():
            if 字段 == "页":
                continue
            try:
                if 字段 == "备注":
                    页.notes_slide.notes_text_frame.text = 值
                    这页成了 = True
                elif 字段 == "标题":
                    形 = 按名字找(页, "页眉标题")
                    if 形 is None:
                        形 = 页.shapes.title
                    if 形 is None:
                        没成.append("第 %d 页找不到标题框，标题没改" % 一条["页"])
                    else:
                        写回(形.text_frame, [值])
                        这页成了 = True
                else:
                    形 = 按名字找(页, 字段)
                    if 形 is None:
                        没成.append("第 %d 页找不到「%s」这一块（这份课件不是新版生成的？）"
                                 "—— 请自己在 WPS 里改" % (一条["页"], 字段))
                    else:
                        写回(形.text_frame, [要点前缀 + x for x in 值] if 字段 == "要点" else [值])
                        这页成了 = True
            except Exception as 异常:
                没成.append("第 %d 页的「%s」没改成（%s）" % (一条["页"], 字段,
                                                     type(异常).__name__))
        if 这页成了:
            好页.append(一条["页"])

    临时 = str(课件路径) + ".改中" + "." + uuid.uuid4().hex[:8]
    演示文稿.save(临时)
    os.replace(临时, str(课件路径))
    if isinstance(警告, list) and 备份名:
        警告.append("动手之前把原样存成了 " + 备份名 + "（在 课件输出/_备份/ 里）")
    return sorted(set(好页)), 没成
