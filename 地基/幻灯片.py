
import os
import uuid

from 地基 import 配置, 课件

页宽 = 13.333
页高 = 7.5
边距 = 0.75
内容宽 = 页宽 - 2 * 边距

主蓝 = (0x1A, 0x3C, 0x6E)
淡蓝 = (0xEC, 0xF1, 0xF8)
强调红 = (0xC0, 0x39, 0x2B)
正文色 = (0x22, 0x22, 0x22)
浅灰 = (0x8A, 0x93, 0xA3)
白 = (0xFF, 0xFF, 0xFF)

字体 = "微软雅黑"

类型们 = ("封面", "知识", "例题", "练习", "小结")

可插图后缀 = (".jpg", ".jpeg", ".png", ".gif")
可配图的类型 = ("知识", "例题")

图左 = 边距 + 6.45
图宽 = 内容宽 - 6.45
图上 = 1.55
图高 = 4.9
文字窄宽 = 6.0


def _警(上下文, 话):
    if isinstance(上下文, dict) and isinstance(上下文.get("警告"), list):
        上下文["警告"].append(话)


def _取图(一页, 上下文):
    名字 = str((一页 or {}).get("图") or "").strip()
    if not 名字:
        return None
    根 = (上下文 or {}).get("图目录") or 配置.取目录("材料")
    真根 = os.path.realpath(根)
    路径 = os.path.realpath(名字 if os.path.isabs(名字) else os.path.join(真根, 名字))
    if not 路径.startswith(真根 + os.sep) or not os.path.isfile(路径):
        _警(上下文, "这一页没配上图（材料/ 里找不到这个文件）：" + 名字[:40])
        return None
    if os.path.splitext(路径)[1].lower() not in 可插图后缀:
        _警(上下文, "这个格式这一版不收（只收 JPEG / PNG / GIF）：" + 名字[:40])
        return None
    return 路径


def _加图(页, 路径):
    from pptx.util import Inches
    形 = 页.shapes.add_picture(路径, Inches(图左), Inches(图上), width=Inches(图宽))
    形.name = "配图"
    高 = 形.height / 914400
    if 高 > 图高:
        形.height = Inches(图高)
        形.width = Inches(形.width / 914400 * (图高 / 高))
        高 = 图高
    形.left = Inches(图左 + 图宽 - 形.width / 914400)
    形.top = Inches(图上 + (图高 - 高) / 2)
    return 形


def _色(元组):
    from pptx.dml.color import RGBColor
    return RGBColor(*元组)


def _加矩形(页, 左, 上, 宽, 高, 颜色, 圆角=False, 线色=None, 名字=""):
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches
    形 = 页.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if 圆角 else MSO_SHAPE.RECTANGLE,
        Inches(左), Inches(上), Inches(宽), Inches(高))
    形.fill.solid()
    形.fill.fore_color.rgb = _色(颜色)
    if 线色 is None:
        形.line.fill.background()
    else:
        形.line.color.rgb = _色(线色)
    形.shadow.inherit = False
    if 名字:
        形.name = 名字
    return 形


def _加文字(页, 左, 上, 宽, 高, 段落们, 字号, 加粗=False, 颜色=正文色,
          行距=1.3, 对齐=None, 段后=8, 名字=""):
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Inches, Pt
    框 = 页.shapes.add_textbox(Inches(左), Inches(上), Inches(宽), Inches(高))
    if 名字:
        框.name = 名字
    文本框 = 框.text_frame
    文本框.word_wrap = True

    for 序号, 一段 in enumerate(段落们):
        段 = 文本框.paragraphs[0] if 序号 == 0 else 文本框.add_paragraph()
        段.text = str(一段)
        段.line_spacing = 行距
        段.space_after = Pt(段后)
        if 对齐 is not None:
            段.alignment = 对齐
        for 片段 in 段.runs:
            片段.font.name = 字体
            片段.font.size = Pt(字号)
            片段.font.bold = 加粗
            片段.font.color.rgb = _色(颜色)
    return 框


def _取要点(一页):
    return 课件.取要点(一页)


def _页眉(页, 一页):
    _加矩形(页, 边距, 0.52, 0.1, 0.66, 主蓝, 名字="标题条")
    _加文字(页, 边距 + 0.3, 0.42, 内容宽 - 0.3, 0.95,
          [一页.get("title", "")], 28, 加粗=True, 颜色=主蓝, 行距=1.0, 段后=0,
          名字="页眉标题")


def _页脚(页, 课题, 序号, 总页数):
    from pptx.util import Inches, Pt
    框 = 页.shapes.add_textbox(Inches(边距), Inches(6.92),
                              Inches(内容宽), Inches(0.34))
    框.name = "页脚"
    文本框 = 框.text_frame
    文本框.word_wrap = False
    段 = 文本框.paragraphs[0]
    段.text = "%s　·　第 %d / %d 页" % (课题, 序号, 总页数)
    for 片段 in 段.runs:
        片段.font.name = 字体
        片段.font.size = Pt(10)
        片段.font.color.rgb = _色(浅灰)


def _排封面(页, 一页, 课题, 序号=None, 总页数=None, 上下文=None):
    from pptx.util import Inches, Pt
    _加矩形(页, 0, 0, 页宽, 页高, 主蓝)
    _加矩形(页, 0, 页高 - 1.05, 页宽, 1.05, (0x14, 0x2F, 0x57))

    标题 = 一页.get("title") or 课题
    _加文字(页, 1.35, 2.25, 页宽 - 2.7, 1.5, [标题], 44, 加粗=True,
          颜色=白, 行距=1.1, 段后=0)
    _加矩形(页, 1.35, 3.95, 1.6, 0.06, 强调红)

    副题 = _取要点(一页)
    if 副题:
        _加文字(页, 1.35, 4.35, 页宽 - 2.7, 1.2, 副题, 18,
              颜色=(0xC9, 0xD6, 0xE8), 行距=1.4, 段后=4)
    _加文字(页, 1.35, 6.15, 页宽 - 2.7, 0.4,
          ["「邵新」辅助教育系统"], 12, 颜色=(0x8F, 0xA8, 0xC8), 段后=0)


def _排知识(页, 一页, 课题, 序号, 总页数, 上下文=None):
    上 = 1.55
    _页眉(页, 一页)
    图路径 = _取图(一页, 上下文)
    宽 = 文字窄宽 if 图路径 else 内容宽 - 0.1

    关键句 = str(一页.get("关键句") or "").strip()
    if 关键句:
        _加矩形(页, 边距, 上, 宽, 1.0, 淡蓝, 圆角=True)
        _加文字(页, 边距 + 0.35, 上 + 0.16, 宽 - 0.7, 0.7,
              [关键句], 21, 加粗=True, 颜色=主蓝, 行距=1.0, 段后=0, 名字="关键句")
        上 += 1.28

    要点 = _取要点(一页)
    if 要点:
        _加文字(页, 边距 + 0.05, 上, 宽, 页高 - 上 - 0.7,
              ["·  " + 一条 for 一条 in 要点], 18 if 图路径 else 20, 行距=1.35, 段后=12,
              名字="要点")
    if 图路径:
        _加图(页, 图路径)
        if isinstance(上下文, dict) and isinstance(上下文.get("配图页"), list):
            上下文["配图页"].append(序号)
    _页脚(页, 课题, 序号, 总页数)


def _排例题(页, 一页, 课题, 序号, 总页数, 上下文=None):
    from pptx.util import Inches, Pt
    _页眉(页, 一页)
    上 = 1.55
    图路径 = _取图(一页, 上下文)
    宽 = 文字窄宽 if 图路径 else 内容宽
    关键句 = str(一页.get("关键句") or "").strip()
    if 关键句:
        _加矩形(页, 边距, 上, 宽, 1.22, 淡蓝, 圆角=True)
        _加矩形(页, 边距, 上, 0.08, 1.22, 强调红)
        _加文字(页, 边距 + 0.4, 上 + 0.22, 宽 - 0.8, 0.85,
              [关键句], 22 if 图路径 else 26, 加粗=True, 颜色=主蓝, 行距=1.0, 段后=0,
              名字="关键句")
        上 += 1.5

    要点 = _取要点(一页)
    if 要点:
        _加文字(页, 边距 + 0.05, 上, 宽, 页高 - 上 - 0.7,
              ["·  " + 一条 for 一条 in 要点], 17 if 图路径 else 19, 行距=1.35, 段后=11,
              名字="要点")
    if 图路径:
        _加图(页, 图路径)
        if isinstance(上下文, dict) and isinstance(上下文.get("配图页"), list):
            上下文["配图页"].append(序号)
    _页脚(页, 课题, 序号, 总页数)


def _排练习(页, 一页, 课题, 序号, 总页数, 上下文=None):
    _页眉(页, 一页)
    上 = 1.6
    关键句 = str(一页.get("关键句") or "").strip()
    if 关键句:
        _加文字(页, 边距 + 0.05, 上, 内容宽 - 0.1, 0.6,
              [关键句], 19, 颜色=强调红, 行距=1.1, 段后=0, 名字="关键句")
        上 += 0.85

    要点 = _取要点(一页)
    if 要点:
        _加文字(页, 边距 + 0.05, 上, 内容宽 - 0.1, 页高 - 上 - 0.7,
              ["%d.  %s" % (号, 一条) for 号, 一条 in enumerate(要点, 1)],
              20, 行距=1.4, 段后=13, 名字="要点")
    _页脚(页, 课题, 序号, 总页数)


def _排小结(页, 一页, 课题, 序号, 总页数, 上下文=None):
    上 = 1.6
    _页眉(页, 一页)
    要点 = _取要点(一页)
    关键句 = str(一页.get("关键句") or "").strip()

    if 要点:
        _加文字(页, 边距 + 0.05, 上, 内容宽 - 0.1, 页高 - 上 - 1.6,
              ["✔  " + 一条 for 一条 in 要点], 20, 行距=1.4, 段后=13,
              名字="要点")

    if 关键句:
        _加矩形(页, 边距, 6.05, 内容宽, 0.72, 主蓝, 圆角=True)
        _加文字(页, 边距 + 0.35, 6.18, 内容宽 - 0.7, 0.5,
              [关键句], 18, 加粗=True, 颜色=白, 行距=1.0, 段后=0, 名字="关键句")
    _页脚(页, 课题, 序号, 总页数)


排页面 = {"封面": _排封面, "知识": _排知识, "例题": _排例题,
         "练习": _排练习, "小结": _排小结}


def 写出(课件数据, 输出路径, 上下文=None):
    try:
        from pptx import Presentation
        from pptx.util import Inches
    except ImportError as 异常:
        raise RuntimeError("没有安装 python-pptx，无法导出 PPTX。安装：pip install python-pptx") from 异常

    课件数据 = 课件数据 if isinstance(课件数据, dict) else {}
    课件.归一类型(课件数据)
    页面们 = [一页 for 一页 in (课件数据.get("pages") or []) if isinstance(一页, dict)]
    课题 = str(课件数据.get("title") or "课件").strip()
    上下文 = 上下文 if isinstance(上下文, dict) else {}
    上下文.setdefault("图目录", 配置.取目录("材料"))
    上下文.setdefault("警告", [])
    上下文.setdefault("配图页", [])

    演示文稿 = Presentation()
    演示文稿.slide_width = Inches(页宽)
    演示文稿.slide_height = Inches(页高)
    空白版式 = 演示文稿.slide_layouts[6]

    总页数 = len(页面们)
    for 序号, 一页 in enumerate(页面们, 1):
        页 = 演示文稿.slides.add_slide(空白版式)
        类型 = str(一页.get("kind") or "知识").strip()
        if 类型 not in 排页面:
            类型 = "知识"
        排页面[类型](页, 一页, 课题, 序号, 总页数, 上下文)
        if 类型 not in 可配图的类型 and str(一页.get("图") or "").strip():
            _警(上下文, "第 %d 页（%s）填了图，已忽略：只有【知识】【例题】两页配图" % (序号, 类型))

        讲稿 = 一页.get("script")
        if 讲稿:
            页.notes_slide.notes_text_frame.text = str(讲稿)

    if not 页面们:
        演示文稿.slides.add_slide(空白版式)

    临时 = str(输出路径) + "." + uuid.uuid4().hex[:8] + ".写出中"
    try:
        演示文稿.save(临时)
        os.replace(临时, 输出路径)
    finally:
        if os.path.isfile(临时):
            try:
                os.remove(临时)
            except OSError:
                pass
    return 输出路径
