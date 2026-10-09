
import hashlib
import uuid
import json
import os
import re
import time


def 抠出JSON对象(原文, 说明="这段回复"):
    文字 = str(原文 or "")
    起点, 终点 = 文字.find("{"), 文字.rfind("}")
    if 起点 < 0 or 终点 <= 起点:
        raise ValueError(说明 + "里没有 JSON 对象")
    try:
        数据 = json.loads(文字[起点:终点 + 1])
    except Exception as 异常:
        raise ValueError(说明 + "里的 JSON 不合法：" + type(异常).__name__) from 异常
    if not isinstance(数据, dict):
        raise ValueError(说明 + "里最外层不是对象")
    return 数据


def 取出JSON(原文):
    说明 = '课件 JSON 结构不对，需要 {"title": ..., "pages": [...]}'
    try:
        数据 = 抠出JSON对象(原文, "课件生成回复")
    except ValueError as 异常:
        raise ValueError(说明) from 异常
    if not 数据.get("pages"):
        raise ValueError(说明)
    return 数据


def 安全文件名(原文):
    结果 = re.sub(r'[\\/:*?"<>|\s]+', "-", str(原文 or "").strip())
    return 结果.strip("-") or "课件"


def 页面清单(数据):
    return [页 for 页 in 数据.get("pages", []) if isinstance(页, dict)]


def 要点成行(要点):
    if isinstance(要点, str):
        return 要点.strip()
    return "\n".join(str(一条) for 一条 in (要点 or []))


def 取要点(一页):
    要点 = (一页 or {}).get("points")
    if isinstance(要点, str):
        return [每一行.strip() for 每一行 in 要点.splitlines() if 每一行.strip()]
    if isinstance(要点, (list, tuple)):
        return [str(一条).strip() for 一条 in 要点 if str(一条).strip()]
    return []


能识别的类型 = ("封面", "知识", "例题", "练习", "小结")
例题味 = ("试一试", "试一", "算一算", "例题", "推导")
练习味 = ("练",)


def 归一类型(数据):
    页面们 = 页面清单(数据)
    总数 = len(页面们)
    for 序号, 页 in enumerate(页面们):
        if str(页.get("kind") or "").strip() in 能识别的类型:
            continue
        标题 = str(页.get("title") or "")
        if 序号 == 0:
            页["kind"] = "封面"
        elif 序号 == 总数 - 1:
            页["kind"] = "小结"
        elif any(词 in 标题 for 词 in 练习味):
            页["kind"] = "练习"
        elif any(词 in 标题 for 词 in 例题味):
            页["kind"] = "例题"
        else:
            页["kind"] = "知识"
    return 数据


def 转文字(数据):
    归一类型(数据)
    行表 = [str(数据.get("title", "课件"))]
    for 序号, 页 in enumerate(页面清单(数据), 1):
        类型 = str(页.get("kind") or "").strip()
        行表.append("")
        行表.append("第 " + str(序号) + " 页"
                  + ("（" + 类型 + "）" if 类型 else "")
                  + "｜" + str(页.get("title", "")))
        if 页.get("关键句"):
            行表.append("★ " + str(页["关键句"]))
        for 要点 in 取要点(页):
            行表.append("- " + 要点)
        if 页.get("script"):
            行表.append("讲稿：" + str(页["script"]))
    return "\n".join(行表)




def 指纹(路径):
    路径 = str(路径)
    if not os.path.isfile(路径):
        raise ValueError("读不到这份课件了（已被删掉或改了名）：" + os.path.basename(路径))
    try:
        尺寸 = os.path.getsize(路径)
        改动时间 = int(os.path.getmtime(路径))
        哈希 = hashlib.sha256()
        with open(路径, "rb") as 文件:
            一段 = 文件.read(1024 * 1024)
            while 一段:
                哈希.update(一段)
                一段 = 文件.read(1024 * 1024)
    except OSError as 异常:
        raise ValueError("这份课件现在读不了（多半是正被 WPS 打开着）：" + os.path.basename(路径)
                       + "。请先把 WPS 里那份关掉，或另存一份再试。") from 异常
    return {"改动时间": 改动时间, "大小": 尺寸, "内容": 哈希.hexdigest()[:16]}


def 指纹一样吗(甲, 乙):
    return bool(甲) and bool(乙) and 甲.get("内容") == 乙.get("内容")


def 读各页(路径):
    try:
        from pptx import Presentation
    except ImportError as 异常:
        raise RuntimeError("没有安装 python-pptx，读不了课件。安装：pip install python-pptx") from 异常

    路径 = str(路径)
    if not os.path.isfile(路径):
        raise ValueError("读不到这份课件了（已被删掉或改了名）：" + os.path.basename(路径))
    try:
        演示文稿 = Presentation(路径)
    except OSError as 异常:
        raise ValueError("这份课件现在打不开（多半是正被 WPS 打开着）：" + os.path.basename(路径)
                       + "。请先把 WPS 里那份关掉再试。") from 异常
    except Exception as 异常:
        raise ValueError("这份文件不像是 .pptx：" + os.path.basename(路径)
                       + "（" + type(异常).__name__ + "）") from 异常

    结果 = []
    for 序号, 页 in enumerate(演示文稿.slides, 1):
        文字们 = []
        for 形 in 页.shapes:
            try:
                if 形.has_text_frame and 形.text_frame.text.strip():
                    文字们.append(形.text_frame.text.strip())
            except Exception:
                continue
        备注 = ""
        try:
            if 页.has_notes_slide:
                备注 = (页.notes_slide.notes_text_frame.text or "").strip()
        except Exception:
            备注 = ""
        标题 = ""
        try:
            if 页.shapes.title is not None and 页.shapes.title.text_frame.text.strip():
                标题 = 页.shapes.title.text_frame.text.strip().splitlines()[0]
        except Exception:
            标题 = ""
        if not 标题 and 文字们:
            标题 = 文字们[0].splitlines()[0].strip()
        按名字 = {}
        for 形 in 页.shapes:
            try:
                if 形.name in ("要点", "关键句") and 形.has_text_frame:
                    按名字[形.name] = 形.text_frame.text.strip()
            except Exception:
                continue
        结果.append({"页": 序号, "标题": 标题, "文字": "\n".join(文字们), "备注": 备注,
                   "要点": 按名字.get("要点", ""), "关键句": 按名字.get("关键句", "")})
    return 结果


def 写备注(路径, 页到文字):
    try:
        from pptx import Presentation
    except ImportError as 异常:
        raise RuntimeError("没有安装 python-pptx，改不了课件。安装：pip install python-pptx") from 异常

    路径 = str(路径)
    if not os.path.isfile(路径):
        raise ValueError("读不到这份课件了（已被删掉或改了名）：" + os.path.basename(路径))
    try:
        演示文稿 = Presentation(路径)
    except OSError as 异常:
        raise ValueError("这份课件现在打不开（多半正被 WPS 打开着）：" + os.path.basename(路径)
                       + "。请先把 WPS 里那份关掉再试。") from 异常
    except Exception as 异常:
        raise ValueError("这份文件不像是 .pptx：" + os.path.basename(路径)
                       + "（" + type(异常).__name__ + "）") from 异常

    要写 = {}
    for 键, 值 in (页到文字 or {}).items():
        try:
            要写[int(键)] = str(值)
        except (TypeError, ValueError):
            raise ValueError("页号得是数字，拿到的是：" + repr(键)) from None

    写好 = 0
    for 序号, 页 in enumerate(演示文稿.slides, 1):
        if 序号 not in 要写:
            continue
        页.notes_slide.notes_text_frame.text = 要写[序号]
        写好 += 1

    临时 = 路径 + ".写出中" + "." + uuid.uuid4().hex[:8]
    try:
        演示文稿.save(临时)
        os.replace(临时, 路径)
    finally:
        if os.path.isfile(临时):
            try:
                os.remove(临时)
            except OSError:
                pass
    return 写好


def 备份成(路径, 标记="老师改过"):
    import shutil
    目录 = os.path.join(os.path.dirname(str(路径)), "_备份")
    os.makedirs(目录, exist_ok=True)
    主干 = os.path.splitext(os.path.basename(str(路径)))[0]
    戳 = time.strftime("%Y%m%d-%H%M")
    序号 = 1
    while True:
        名 = "%s.%s.%s%s.pptx" % (主干, 标记, 戳, "" if 序号 == 1 else "-" + str(序号))
        if not os.path.isfile(os.path.join(目录, 名)):
            break
        序号 += 1
    shutil.copy2(str(路径), os.path.join(目录, 名))
    return 名


def 文本指纹(文字):
    return hashlib.sha256(str(文字 or "").encode("utf-8")).hexdigest()[:16]


def 备注指纹(一页):
    return 文本指纹((一页 or {}).get("备注"))
