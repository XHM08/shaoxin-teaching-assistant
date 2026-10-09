
import html as html模块
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import urllib.parse

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from 地基 import 大模型, 配置, 课件, 课件图
from 能力 import 讲稿配音, 用已有课件

可能的Edge = [
    r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe",
    r"C:\Program Files\Microsoft\Edge\Application\msedge.exe",
]
标记 = "自检结果标记"
结果 = []


def 检查(名字, 通过, 说明=""):
    结果.append((名字, bool(通过), 说明))


def 找Edge():
    for 路径 in 可能的Edge:
        if os.path.isfile(路径):
            return 路径
    return None


def 后端():
    输出 = 配置.取目录("课件输出")
    主干 = "外来课件-自检-%d" % os.getpid()
    课名 = 主干 + ".pptx"
    路径 = os.path.join(输出, 课名)
    备份目录 = os.path.join(输出, "_备份")
    备份_前 = set(os.listdir(备份目录)) if os.path.isdir(备份目录) else set()
    要点甲, 要点乙 = "· 甲条", "· 乙条"
    原有 = {2: "（原有讲稿，别动我）老师自己写的第二页。", 4: "（原有讲稿，别动我）第四页。"}
    内容 = [(1, "地球上的水", "海洋、湖泊、大气里都有水。"), (2, "蒸发", "太阳一晒，水变成水蒸气。"),
          (3, "凝结与降水", "水蒸气遇冷变成小水滴落下来。"), (4, "小结", "蒸发→凝结→降水。")]
    应补 = [号 for 号, _, _ in 内容 if 号 not in 原有]
    新备份 = []

    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    for 号, 标, 文 in 内容:
        页 = prs.slides.add_slide(prs.slide_layouts[6])
        for 上, 字 in ((1.0, 标), (2.2, 文)):
            框 = 页.shapes.add_textbox(Inches(0.8), Inches(上), Inches(8), Inches(1))
            框.text_frame.text = 字
        if 号 == 1:
            要点框 = 页.shapes.add_textbox(Inches(0.8), Inches(3.4), Inches(8), Inches(1.2))
            要点框.name = "要点"
            要点框.text_frame.text = 要点甲
            要点框.text_frame.add_paragraph().text = 要点乙
        if 号 in 原有:
            页.notes_slide.notes_text_frame.text = 原有[号]
    prs.save(路径)
    检查("造出了一份「外来」课件（空白版式，没有我们那套形状名）", True, 课名)

    调过, 配过 = [], []
    原问, 原导出, 原合成 = 大模型.问, 课件图.导出各页, 讲稿配音.合成整份

    def 假问(原文, **kw):
        调过.append(原文)
        号 = next((号 for 号, _, 文 in 内容 if 文 in 原文), "?")
        return "（自检替身）第 %s 页的讲稿。" % 号

    def 假导出(路径参数, 警告=None):
        return {"目录名": 主干 + "_页图", "目录": os.path.join(输出, 主干 + "_页图"),
              "图": [], "页数": len(内容)}

    def 假合成(*a, **k):
        配过.append(1)
        return "", []

    大模型.问, 课件图.导出各页, 讲稿配音.合成整份 = 假问, 假导出, 假合成
    try:
        回 = 用已有课件.处理(课名, "demo-real", "不")
    finally:
        大模型.问, 课件图.导出各页, 讲稿配音.合成整份 = 原问, 原导出, 原合成
        if os.path.isdir(备份目录):
            新备份 = sorted(set(os.listdir(备份目录)) - 备份_前)

    检查("要点成行：字符串原样、列表连成多行、空的安全",
       课件.要点成行("· 甲\n· 乙") == "· 甲\n· 乙"
       and 课件.要点成行(["· 甲", "· 乙"]) == "· 甲\n· 乙"
       and 课件.要点成行(None) == "" and 课件.要点成行([]) == "",
       "串=ok 列表=ok 空=ok")

    页们 = {一["页"]: 一 for 一 in 课件.读各页(路径)}
    检查("只给没有讲稿的页补（原有的备注一个字没动）",
       页们[2]["备注"] == 原有[2] and 页们[4]["备注"] == 原有[4]
       and all(页们[号]["备注"].strip() for 号 in 应补),
       "第 2/4 页 = %s / %s" % (页们[2]["备注"][:10], 页们[4]["备注"][:10]))
    检查("标题与正文也没被动（承诺是只补讲稿）",
       all(页们[号]["标题"].strip() == 标 and 文 in 页们[号]["文字"] for 号, 标, 文 in 内容))
    检查("提示里明说补了几页、并明说原有的没动",
       ("补了 %d 页讲稿" % len(应补)) in 回["提示"] and "一个字没动" in 回["提示"],
       str(回["提示"])[:80])
    _要点 = str(回["条目"][0].get("要点") or "")
    检查("要点是两行、不是被逐字拆开的（真课件里的要点是字符串，模型给的是列表）",
       ("甲条" in _要点 and "乙条" in _要点 and _要点.count(chr(10)) == 1
        and len(_要点) < 40),
       "条目[0]的要点 = %r" % _要点[:60])

    检查("键集与播放器要的那 6 个字段一致（上课那一屏不用改）",
       set(回["条目"][0].keys()) == {"页", "标题", "要点", "讲稿", "音频", "页图"},
       str(sorted(回["条目"][0].keys())))
    每页次数 = {号: sum(1 for x in 调过 if 文 in x) for 号, _, 文 in 内容 if 号 in 应补}
    检查("模型调用次数 = 缺讲稿的页数，且每页只调一次",
       len(调过) == len(应补) and set(每页次数.values()) == {1} and len(每页次数) == len(应补),
       "调了 %d 次，应为 %d；每页次数 %s" % (len(调过), len(应补), 每页次数))
    检查("每页图都拿到了地址（播放页要显示他的样子）",
       all(x.get("页图") for x in 回["条目"]))
    检查("填了「不」就真的一次都没配音（替身数为 0）", not 配过, "配音被调了 %d 次" % len(配过))
    本次 = [名 for 名 in 新备份 if 主干 in 名]
    检查("动的之前留了备份", bool(本次), str(新备份[:1]))

    主干2 = "自检-半途失败-%d" % os.getpid()
    自建2 = os.path.join(输出, 主干2 + ".pptx")
    备份_前2 = set(os.listdir(备份目录)) if os.path.isdir(备份目录) else set()
    prs2 = Presentation()
    for 正文2 in ("甲页正文", "乙页正文"):
        页2 = prs2.slides.add_slide(prs2.slide_layouts[6])
        框2 = 页2.shapes.add_textbox(Inches(0.8), Inches(1), Inches(8), Inches(1))
        框2.text_frame.text = 正文2
    prs2.save(自建2)

    def 会炸的问(原文, *a, **k):
        if "乙页正文" in 原文:
            raise RuntimeError("假装模型在第 2 页挂了")
        return "（自检替身）第 1 页的讲稿。"

    大模型.问 = 会炸的问
    try:
        回2 = 用已有课件.处理(os.path.basename(自建2), "demo-real", "不")
    finally:
        大模型.问 = 原问
    页们2 = {一["页"]: 一 for 一 in 课件.读各页(自建2)}
    检查("一页失败不回滚：已写好的那页讲稿留下、失败那页点名",
       页们2[1]["备注"].strip() == "（自检替身）第 1 页的讲稿。"
       and not 页们2[2]["备注"].strip()
       and "第 2 页" in str(回2.get("提示")),
       "第1页=%r 第2页=%r 提示=%s" % (页们2[1]["备注"][:16], 页们2[2]["备注"][:16], str(回2.get("提示"))[:60]))
    for 目标2 in (自建2, os.path.join(输出, 主干2 + "_页图"), os.path.join(输出, 主干2 + "_配音")):
        if os.path.isdir(目标2):
            shutil.rmtree(目标2, ignore_errors=True)
        elif os.path.isfile(目标2):
            try:
                os.remove(目标2)
            except OSError:
                pass
    if os.path.isdir(备份目录):
        for 名2 in sorted(set(os.listdir(备份目录)) - 备份_前2):
            if 主干2 in 名2:
                try:
                    os.remove(os.path.join(备份目录, 名2))
                except OSError:
                    pass
    if 本次:
        备表 = {一["页"]: 一 for 一 in 课件.读各页(os.path.join(备份目录, 本次[0]))}
        检查("备份是可用的原样（第 2/4 页备注在、第 1/3 页仍是空）",
           备表[2]["备注"] == 原有[2] and 备表[4]["备注"] == 原有[4]
           and not 备表[1]["备注"].strip() and not 备表[3]["备注"].strip())

    return 课名, 输出, 主干, 路径, 备份目录, 本次


def 清理后端产物(输出, 主干, 路径, 备份目录, 本次):
    for 目标 in (路径, os.path.join(输出, 主干 + "_页图"), os.path.join(输出, 主干 + "_配音")):
        if os.path.isdir(目标):
            shutil.rmtree(目标, ignore_errors=True)
        elif os.path.isfile(目标):
            try:
                os.remove(目标)
            except OSError:
                pass
    for 名 in 本次:
        try:
            os.remove(os.path.join(备份目录, 名))
        except OSError:
            pass


测试脚本 = """
<script>
(async () => {
  const 结果 = [];
  const 记 = (名, 过, 说) => 结果.push({名前: 名, 过: !!过, 说: String(说 || '')});
  const 等 = (毫秒) => new Promise(r => setTimeout(r, 毫秒));
  await 等(2500);                       // 页面自己的初始化 IIFE 是异步的
  try {

  const 课名 = __课名__;
  const 调用 = [];
  window.run = async (能力, 参数) => {   // 替身：不花钱、不等 COM，只记录"调了什么、带了什么参数"
    调用.push([能力, 参数]);
    return { 提示: '（替身）用你这份课件（共 2 页）：补了 1 页讲稿，另外 1 页本来就有讲稿（一个字没动）',
             课名: 参数.课件文件 || '',
             条目: [{页: 1, 标题: '地球上的水', 要点: '', 讲稿: '甲', 音频: '',
                   页图: '/课件图/自检页图_页图/page1.png'},
                   {页: 2, 标题: '蒸发', 要点: '', 讲稿: '乙', 音频: '',
                   页图: '/课件图/自检页图_页图/page2.png'}] };
  };

  await show('联动教学');
  await 等(800);
  const 源 = $('lecSource');
  记('① 界面上有「这一课怎么来」这个选择（两条路）', !!源 && 源.options.length === 2,
     源 ? Array.from(源.options).map(o => o.value).join('/') : '（没有这个下拉）');

  源.value = '已有';
  源.dispatchEvent(new Event('change'));
  await 等(1200);                        // 等课件下拉真去取一次列表
  记('② 切到「用已有课件」：课件框出现、课题与页数收起',
     !$('lecDeckBox').classList.contains('hide') && $('lecTopicBox').classList.contains('hide')
     && $('lecSlidesBox').classList.contains('hide'),
     '课件框hide=' + $('lecDeckBox').classList.contains('hide')
     + ' 课题hide=' + $('lecTopicBox').classList.contains('hide'));
  const 选项 = Array.from($('lecDeck').options).map(o => o.value);
  记('③ 课件下拉真列出来了（不是空、也不是占位句）', 选项.length > 0 && !!选项[0], 选项.slice(0, 3).join(' | '));
  记('④ 刚放进 课件输出/ 的那份就在列表里（新的排最前）', 选项[0] === 课名, 选项.slice(0, 2).join(' | '));

  $('lecDeck').value = 课名;
  const 技能选项 = Array.from($('lecSkill').options).map(o => o.value).filter(Boolean);
  记('⑤ 技能包下拉也填好了（讲稿要按某位老师的讲法写）', 技能选项.length > 0, 技能选项.slice(0, 3).join(' | '));
  $('lecSkill').value = 技能选项[0];
  $('lecWantVoice').checked = false;
  $('btnLecStart').click();
  for (let i = 0; i < 40 && !LEC.items.length; i++) { await 等(200); }

  const 参数 = 调用.length ? 调用[0][1] : {};
  记('⑥ 点「开始上课」调的是「用已有课件」', 调用.length === 1 && 调用[0][0] === '用已有课件',
     调用.length ? 调用[0][0] : '（一次都没调）');
  // 比集合，不比 sort 后的字符串：JS 排中文按 UTF-16 码位，不是拼音序，'课件文件' 不在最前
  const 参数键 = Object.keys(参数);
  记('⑦ 参数只有课件/技能包/要不要配音（没有课题、页数）',
     参数键.length === 3 && ['课件文件', '技能包', '要不要配音'].every(k => 参数键.indexOf(k) >= 0)
     && 参数.课件文件 === 课名 && 参数.要不要配音 === '不',
     JSON.stringify(参数));
  记('⑧ 播放器出来了、标题带课名', !$('lecPlayer').classList.contains('hide')
     && $('lecTitle').textContent.indexOf(课名.replace('.pptx', '')) >= 0,
     $('lecTitle').textContent);
  记('⑨ 第一页显示的是他的样子（页图 src 设上了）',
     ($('lecShot').getAttribute('src') || '').indexOf('page1.png') >= 0,
     $('lecShot').getAttribute('src') || '（没设）');
  记('⑩ 没音频时音频条收起（不给老师一个空播放器）', $('lecVoiceBox').classList.contains('hide'));
  记('⑪ 提示里写清了补了几页讲稿', $('lecOutMsg').textContent.indexOf('补了 1 页讲稿') >= 0,
     $('lecOutMsg').textContent.slice(0, 60));

  $('btnLecStop').click();
  await 等(400);
  记('⑫ 结束上课后：面板复位（回到「生成」那一档，课题框回来）',
     源.value === '生成' && !$('lecTopicBox').classList.contains('hide')
     && $('lecDeckBox').classList.contains('hide'),
     '源=' + 源.value + ' 课题hide=' + $('lecTopicBox').classList.contains('hide'));
  源.value = '已有';
  源.dispatchEvent(new Event('change'));
  await 等(300);
  记('⑬ 再切一次仍然正常（课件框回来、课题收起）',
     !$('lecDeckBox').classList.contains('hide') && $('lecTopicBox').classList.contains('hide'));

  } catch (错) {
    // 脚本自己抛错也要留下痕迹：不然外面只看到“没找到标记”，跟“页面没起来”混成一谈
    结果.push({名前: '自检脚本自己没跑完（看这条的原文）', 过: false,
             说: String((错 && 错.message) || 错).slice(0, 160)});
  }
  const 盒 = document.createElement('div');
  盒.id = '__标记__';
  盒.textContent = JSON.stringify(结果);
  盒.style.display = 'none';
  document.body.appendChild(盒);
})();
</script>
"""


def 界面(课名):
    edge = 找Edge()
    if not edge:
        检查("界面那半（需要无头 Edge）", True, "跳过：这台机器上找不到 Edge，这一条没验")
        return "这台机器上没有无头 Edge，界面那半一项都没验"
    网页目录 = os.path.join(配置.根目录, "网页")
    原页 = os.path.join(网页目录, "首页.html")
    临时页 = os.path.join(网页目录, "_用已有课件自检_%d.html" % os.getpid())
    端口 = 配置.取端口()
    try:
        with open(原页, encoding="utf-8") as 文件:
            页面 = 文件.read()
    except OSError as 异常:
        检查("读得到 首页.html", False, type(异常).__name__)
        return
    脚本 = 测试脚本.replace("__课名__", json.dumps(课名)).replace("__标记__", 标记)
    with open(临时页, "w", encoding="utf-8") as 文件:
        文件.write(页面.replace("</body>", 脚本 + "</body>", 1))
    地址 = "http://127.0.0.1:%d%s" % (端口, urllib.parse.quote("/网页/" + os.path.basename(临时页)))
    try:
        跑 = subprocess.run([edge, "--headless=new", "--disable-gpu", "--no-first-run",
                           "--disable-extensions", "--virtual-time-budget=20000",
                           "--dump-dom", 地址],
                          capture_output=True, text=True, encoding="utf-8", timeout=180)
        dom = 跑.stdout or ""
        找 = re.search(r'id="' + 标记 + r'"[^>]*>(.*?)</div>', dom, re.S)
        if not 找:
            检查("界面上那半跑到结尾（页面能起、脚本没抛错）", False,
               "没找到结果标记；DOM %d 字符。可能是服务没起或页面脚本报错" % len(dom))
        else:
            for 一 in json.loads(html模块.unescape(找.group(1))):
                检查(一["名前"], 一["过"], 一["说"])
    except subprocess.TimeoutExpired:
        检查("界面那半跑得完", False, "超时（180 秒）")
    except Exception as 异常:
        检查("界面那半跑得完", False, type(异常).__name__ + "：" + str(异常)[:80])
    finally:
        try:
            os.remove(临时页)
        except OSError:
            pass


def main():
    端口 = 配置.取端口()
    import socket
    探 = socket.socket()
    探.settimeout(3)
    try:
        探.connect(("127.0.0.1", 端口))
    except Exception:
        print("  服务没起（127.0.0.1:%d 连不上）。界面那半要真取一次 /api/能力 才有课件下拉。" % 端口)
        print("  先启动服务：双击 启动.bat，或 python 本地服务.py")
        return 1
    finally:
        探.close()

    主干兜底 = "外来课件-自检-%d" % os.getpid()
    输出兜底 = 配置.取目录("课件输出")
    课名 = 输出 = 主干 = 路径 = 备份目录 = 本次 = None
    跳过的 = None
    try:
        课名, 输出, 主干, 路径, 备份目录, 本次 = 后端()
        跳过的 = 界面(课名)
    finally:
        真输出 = 输出 or 输出兜底
        真主干 = 主干 or 主干兜底
        清理后端产物(真输出, 真主干,
                   路径 or os.path.join(真输出, 真主干 + ".pptx"),
                   备份目录 or os.path.join(真输出, "_备份"),
                   本次 or [])

    print()
    坏 = 0
    for 名字, 过, 说明 in 结果:
        print(("  通过  " if 过 else "  失败  ") + 名字 + ("　→　" + 说明 if 说明 else ""))
        坏 += 0 if 过 else 1
    if 坏:
        print("%d 项失败" % 坏)
        return 1
    if 跳过的:
        print("全部 %d 项通过，但有一半没验（%s）" % (len(结果), 跳过的))
        print("按约定返回 3 = 跳过：装好无头 Edge 再跑这一条")
        return 3
    print("全部 %d 项通过" % len(结果))
    return 0


if __name__ == "__main__":
    sys.exit(main())
