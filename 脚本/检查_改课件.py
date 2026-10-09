
import io
import os
import shutil
import sys
import tempfile
import wave

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from 地基 import 课件, 课件改, 配置, 语音
from 能力 import _课件, 讲稿配音, 改课件

原取目录 = 配置.取目录
原问 = 课件改.大模型.问
原能配音 = 讲稿配音.能配音吗
原合成 = 语音.合成
临时 = tempfile.mkdtemp(prefix="改课件自检_")
输出目录 = os.path.join(临时, "课件输出")
os.makedirs(输出目录)
结果 = []
合成过 = []


def 检查(名字, 通过, 说明=""):
    结果.append((名字, bool(通过), 说明))


配置.取目录 = lambda 名: 输出目录 if 名 == "课件输出" else 原取目录(名)


def 备份们(目录):
    d = os.path.join(目录, "_备份")
    return sorted(os.listdir(d)) if os.path.isdir(d) else []


def 备份里有原讲稿(目录, 前缀, 片段):
    import pptx as _pptx
    for 名 in 备份们(目录):
        if 前缀 in 名:
            try:
                演示 = _pptx.Presentation(os.path.join(目录, "_备份", 名))
                for 页 in 演示.slides:
                    备注 = 页.notes_slide.notes_text_frame.text if 页.has_notes_slide else ""
                    if 片段 in 备注:
                        return True
            except Exception:
                continue
    return False


def 造课件():
    数据 = {"title": "改课件夹具", "pages": [
        {"kind": "封面", "title": "封面", "points": ["六年级"], "script": "开场的话"},
        {"kind": "知识", "title": "倒数是什么", "关键句": "分子分母交换位置",
         "points": ["交换分子分母", "零没有倒数"], "script": "原来的讲稿：我们看三分之二。"},
        {"kind": "小结", "title": "小结", "关键句": "记住这一句", "points": ["一条"], "script": "收尾"},
    ]}
    return os.path.join(输出目录, _课件.保存pptx(数据, "自检包", "改课件", _课件.上下文表()))


def 假合成(文本, 参考音, 参考音文本, **其余):
    合成过.append(文本)
    缓冲 = io.BytesIO()
    with wave.open(缓冲, "wb") as 音频:
        音频.setnchannels(1)
        音频.setsampwidth(2)
        音频.setframerate(8000)
        音频.writeframes(b"\x00\x00" * 80)
    return 缓冲.getvalue()


语音.合成 = 假合成
讲稿配音.能配音吗 = lambda: (True, "")
假回复 = """{"改动": [
  {"页": 2, "备注": "老师说要改成口语一点：我们看三分之二。", "要点": ["交换分子分母", "零没有倒数，要记住"], "关键句": "把分子分母换个位置"},
  {"页": 9, "备注": "这一页不存在"},
  {"页": 1, "字号": 40}
], "做不到的": ["第 3 页想加一张图，做不到"]}"""
课件改.大模型.问 = lambda *a, **k: 假回复

路径 = 造课件()
件名 = os.path.basename(路径)
try:
    前印 = 课件.指纹(路径)
    回执 = 改课件.处理(件名, "第 2 页讲得太书面了，改成我平时说话的样子")
    检查("① 只出清单时提示里说的是\"再点一次确认\"，而不是\"改好了\"",
        "再点一次" in 回执["提示"] and "改好了" not in 回执["提示"], 回执["提示"][:56])
    检查("② **没确认就一个字都不改**（文件指纹没动）",
        课件.指纹一样吗(前印, 课件.指纹(路径)),
        "前 %s 后 %s" % (前印["内容"], 课件.指纹(路径)["内容"]))
    检查("③ 清单里越界页号（第 9 页）进了\"做不到的\"，没当成改动",
        any("第 9 页不存在" in x for x in 回执["清单"]["做不到的"])
        and 9 not in [x["页"] for x in 回执["清单"]["改动"]],
        str(回执["清单"]["做不到的"])[:70])
    检查("④ 清单里这一版改不了的字段（字号）也进了\"做不到的\"",
        any("字号" in x for x in 回执["清单"]["做不到的"]), str(回执["清单"]["做不到的"])[:70])

    讲稿配音.合成整份(路径, 1.0)
    合成过.clear()
    回执 = 改课件.处理(件名, "第 2 页讲得太书面了", "好")
    现在 = 课件.读各页(路径)
    检查("⑤ 确认后：备注（讲稿）真的换了",
        "口语一点" in 现在[1]["备注"], 现在[1]["备注"][:26])
    检查("⑥ 确认后：要点按形状名找到并换了（两条，带项目符号）",
        "零没有倒数，要记住" in 现在[1]["要点"] and "·" in 现在[1]["要点"],
        现在[1]["要点"].replace(chr(10), " / ")[:44])
    检查("⑦ 确认后：关键句也换了",
        "把分子分母换个位置" in 现在[1]["关键句"], 现在[1]["关键句"][:26])
    检查("⑧ 动手前留了备份，且是**改之前**那一版（原讲稿还在里面）",
        备份里有原讲稿(输出目录, "自检包-改课件", "原来的讲稿"),
        "备份：" + str(备份们(输出目录)))
    检查("⑨ 改完只重配了改过那页的配音（不是整课重来）",
        len(合成过) == 1 and "口语一点" in 合成过[0], "合成 %d 次：%s" % (len(合成过), 合成过[:1]))

    老路径 = os.path.join(输出目录, "老课件.pptx")
    from pptx import Presentation
    演示 = Presentation(路径)
    for 页 in 演示.slides:
        for 形 in 页.shapes:
            if 形.name in ("要点", "关键句", "页眉标题", "标题条"):
                形.name = "旧形状"
    演示.save(老路径)
    课件改.大模型.问 = lambda *a, **k: '{"改动": [{"页": 2, "要点": ["新的要点"]}], "做不到的": []}'
    回执 = 改课件.处理("老课件.pptx", "第 2 页要点换一下", "好")
    检查("⑩ 找不到「要点」这一块时：明说改不了、请在 WPS 里改（不硬改、不静默）",
        "找不到" in 回执["提示"] and "WPS" in 回执["提示"], 回执["提示"][:70])

    课件改.大模型.问 = lambda *a, **k: '```json\n{"改动": [{"页": 3, "备注": "新的收尾"}], "做不到的": []}\n```'
    清单 = 课件改.出清单(路径, "第 3 页换个收尾")
    检查("⑪ 回复带 ```json 围栏也能解析", len(清单["改动"]) == 1 and 清单["改动"][0]["页"] == 3,
        str(清单)[:70])

    课件改.大模型.问 = lambda *a, **k: "我不太确定你想改哪儿。"
    try:
        课件改.出清单(路径, "随便改改")
        检查("⑫ 模型回垃圾 → 给人话", False, "居然没报错")
    except ValueError as 异常:
        检查("⑫ 模型回垃圾 → 给人话（不是 traceback）", "不是 JSON" in str(异常), str(异常)[:56])
finally:
    配置.取目录 = 原取目录
    课件改.大模型.问 = 原问
    讲稿配音.能配音吗 = 原能配音
    语音.合成 = 原合成
from pptx import Presentation
from pptx.util import Inches

自建 = os.path.join(配置.取目录("课件输出"), "自检-复用回归-%d.pptx" % os.getpid())
prs = Presentation()
for _ in range(2):
    页 = prs.slides.add_slide(prs.slide_layouts[6])
    框 = 页.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
    框.text_frame.text = "复用回归页"
prs.save(自建)
调过 = []


def _数着的假问(原文, *a, **k):
    调过.append(1)
    return '{"改动": [{"页": 2, "备注": "复用测试"}], "做不到的": []}'


课件改.大模型.问 = _数着的假问
try:
    _一 = 改课件.处理(os.path.basename(自建), "把第 2 页备注换成复用测试")
    _二 = 改课件.处理(os.path.basename(自建), "把第 2 页备注换成复用测试", "好")
    检查("确认那一步不再调模型（老师点头的就是落地的那份）", len(调过) == 1,
       "模型被调了 %d 次" % len(调过))
finally:
    课件改.大模型.问 = 原问
    try:
        os.remove(自建)
    except OSError:
        pass
    备份目录 = os.path.join(配置.取目录("课件输出"), "_备份")
    if os.path.isdir(备份目录):
        for 名 in os.listdir(备份目录):
            if "自检-复用回归-%d" % os.getpid() in 名:
                try:
                    os.remove(os.path.join(备份目录, 名))
                except OSError:
                    pass

    shutil.rmtree(临时, ignore_errors=True)

for 序号, (名字, 通过, 说明) in enumerate(结果, 1):
    print("  %2d. %s %s —— %s" % (序号, "通过" if 通过 else "不过", 名字, 说明))

通过数 = sum(1 for _名, 通, _说 in 结果 if 通)
print()
print("通过 %d / 共 %d" % (通过数, len(结果)))
sys.exit(0 if 通过数 == len(结果) else 1)
