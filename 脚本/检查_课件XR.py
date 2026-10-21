
import os
import shutil
import sys
import tempfile

根 = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, 根)

失败 = []
条数 = 0


def 记(名, 通过, 说明=""):
    global 条数
    条数 += 1
    print(("  ✓ " if 通过 else "  ✗ ") + 名 + ("" if 通过 else "　→　" + str(说明)))
    if not 通过:
        失败.append(名)


def 读XR(路径):
    from pptx import Presentation
    演示文稿 = Presentation(路径)
    return len(演示文稿.slides), (演示文稿.core_properties.comments or "").strip()


def 跑检查():
    from 地基 import 幻灯片

    临时 = tempfile.mkdtemp(prefix="课件XR-")
    try:
        数据 = {"title": "自检用课件", "pages": [
            {"kind": "封面", "title": "自检用课件", "points": ["五年级 · 数学"], "script": "开场。", "图": ""},
            {"kind": "小结", "title": "小结", "points": ["记住这一句"], "script": "收尾。", "图": ""},
        ]}
        正本 = os.path.join(临时, "带XR.pptx")
        幻灯片.写出(数据, 正本, {"图目录": 临时})
        页数, 标记 = 读XR(正本)
        记("① 课件写出来了，页数也对", os.path.isfile(正本) and 页数 == 2, "页数 %s" % 页数)
        记("② 文件属性里带着 XR: C", 标记 == "XR: C", "读到的是 %r" % (标记,))

        from pptx import Presentation
        对照 = os.path.join(临时, "无XR.pptx")
        空 = Presentation()
        空.slides.add_slide(空.slide_layouts[6])
        空.save(对照)
        _, 对照标记 = 读XR(对照)


        记("③ 反向验：不经过我们的写出，就没有我们的标记",
           对照标记 != "XR: C", "对照组是 %r" % (对照标记,))
    finally:
        shutil.rmtree(临时, ignore_errors=True)


def 主():
    global 条数
    try:
        跑检查()
    except Exception as 错:

        记("检查本身跑完了", False, "%s: %s" % (type(错).__name__, 错))
    print()
    print("全部 %d 项通过" % 条数 if not 失败 else "通过 %d / 共 %d" % (条数 - len(失败), 条数))
    return 0 if not 失败 else 1


if __name__ == "__main__":
    sys.exit(主())
