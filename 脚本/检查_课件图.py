
import os
import shutil
import sys
import time

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from 地基 import 课件, 课件图, 配置

输出目录 = 配置.取目录("课件输出")
真课件 = os.path.join(输出目录, "demo-real-分数的除法.pptx")
副本 = os.path.join(输出目录, "自检用-页图副本.pptx")
结果 = []


def 检查(名字, 通过, 说明=""):
    结果.append((名字, bool(通过), 说明))


print("  导出用的程序：" + str(课件图.找放映程序() or "（这台机器上没找到 WPS 演示 / PowerPoint）"))

能, 为什么 = 课件图.能导出吗()
if not 能:
    print("  跳过：" + 为什么)
    print("\n0 项检查，1 项跳过（没有放映软件，功能不受影响，只是播放页用文字版）")
    sys.exit(3)

if not os.path.isfile(真课件):
    print("  没有可用的课件做夹具：" + 真课件)
    sys.exit(2)

shutil.copyfile(真课件, 副本)
位置 = 课件图.取位置(副本)
目录名, 一个 = 位置["目录名"], 位置["目录"]
shutil.rmtree(一个, ignore_errors=True)

try:
    警告 = []
    开始 = time.time()
    位置 = 课件图.导出各页(副本, 警告)
    目录名, 图们 = 位置["目录名"], 位置["图"]
    耗时 = time.time() - 开始
    页数 = len(课件.读各页(副本))
    检查("① 导出的张数 = 课件页数（%d 页）" % 页数, len(图们) == 页数,
        "导了 %d 张，用时 %.1f 秒" % (len(图们), 耗时))
    头 = b""
    第一个 = os.path.join(一个, 图们[0]) if 图们 else ""
    if 第一个 and os.path.isfile(第一个):
        with open(第一个, "rb") as 文件:
            头 = 文件.read(8)
    检查("② 导出来的是合法 PNG（魔数 \x89PNG）", 头.startswith(b"\x89PNG"),
        repr(头) + "，%.0f KB" % (os.path.getsize(第一个) / 1024 if 第一个 and os.path.isfile(第一个) else 0))
    检查("③ 台账落盘（下次才知道要不要重导）",
        os.path.isfile(os.path.join(一个, 课件图.台账名)))
    检查("④ 首次导出说的是「已导出」，**不能说「课件变了」**（那是假话）",
        any("已导出" in 一句 for 一句 in 警告) and not any("课件变了" in 一句 for 一句 in 警告),
        "；".join(警告)[:80])

    时间甲 = os.path.getmtime(第一个)
    time.sleep(1.1)
    警告 = []
    课件图.导出各页(副本, 警告)
    检查("⑤ 课件没变 → 不重导（图的原样时间没动）",
        os.path.getmtime(第一个) == 时间甲 and not 警告, "；".join(警告)[:60])

    from pptx import Presentation
    演示 = Presentation(副本)
    list(演示.slides)[0].notes_slide.notes_text_frame.text = "改一下备注，让指纹变化"
    演示.save(副本)
    time.sleep(1.1)
    警告 = []
    课件图.导出各页(副本, 警告)
    检查("⑥ 课件变了 → 重导，而且这时才该说「课件改了」",
        os.path.getmtime(第一个) > 时间甲 and any("课件改了" in 一句 for 一句 in 警告),
        "；".join(警告)[:80])

    检查("⑦ 页图地址形状对（/课件图/<目录>/pageN.png，未编码）",
        课件图.页图地址(目录名, 2) == "/课件图/" + 目录名 + "/page2.png",
        课件图.页图地址(目录名, 2))
finally:
    shutil.rmtree(一个, ignore_errors=True)
    if os.path.isfile(副本):
        os.remove(副本)

print()
坏 = 0
for 名字, 过, 说明 in 结果:
    print(("  通过  " if 过 else "  失败  ") + 名字 + ("　→　" + 说明 if 说明 else ""))
    坏 += 0 if 过 else 1
print("全部 %d 项通过" % len(结果) if not 坏 else "%d 项失败" % 坏)
sys.exit(0 if not 坏 else 1)
