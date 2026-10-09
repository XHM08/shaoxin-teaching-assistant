
import os
import shutil
import sys
import tempfile

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from 地基 import 课件, 配置
from 能力 import _课件

原取目录 = 配置.取目录
临时 = tempfile.mkdtemp(prefix="覆盖保护自检_")
输出目录 = os.path.join(临时, "课件输出")
os.makedirs(输出目录)
结果 = []


def 检查(名字, 通过, 说明=""):
    结果.append((名字, bool(通过), 说明))


配置.取目录 = lambda 名: 输出目录 if 名 == "课件输出" else 原取目录(名)


def 造数据(标记):
    return {"title": "覆盖保护夹具", "pages": [
        {"kind": "封面", "title": "封面", "points": ["六年级"], "script": "开场"},
        {"kind": "知识", "title": "第一页", "关键句": "记号 " + 标记,
         "points": ["一条"], "script": "讲稿 " + 标记},
    ]}


def 生成一次(标记):
    上下文 = _课件.上下文表()
    文件名 = _课件.保存pptx(造数据(标记), "自检包", "覆盖保护", 上下文)
    return os.path.join(输出目录, 文件名), 上下文["警告"]


def 备份们():
    目录 = os.path.join(输出目录, "_备份")
    return sorted(os.listdir(目录)) if os.path.isdir(目录) else []


def 改备注(路径, 文字):
    from pptx import Presentation
    演示 = Presentation(路径)
    list(演示.slides)[1].notes_slide.notes_text_frame.text = 文字
    演示.save(路径)


try:
    路径, 警告 = 生成一次("第一次")
    检查("① 首次生成：不产生备份（本来就没有可覆盖的东西）", 备份们() == [], str(备份们()))
    检查("①b 生成记录写好了（下次才分得清是不是我们生成的）",
        os.path.isfile(_课件.生成记录路径(路径)))

    路径, 警告 = 生成一次("第二次")
    检查("② 没人改过就再生成 → 不产生备份、也不谎报\"你改过\"",
        备份们() == [] and not any("不像是我们刚生成的那版" in 一句 for 一句 in 警告),
        "备份=%s 警告=%s" % (备份们(), 警告))

    改备注(路径, "老师自己改的讲稿【探针】")
    路径, 警告 = 生成一次("第三次")
    备份 = 备份们()
    检查("③ 改过之后再生成 → 先备份，且提醒里带上备份名",
        len(备份) == 1 and any("不像是我们刚生成的那版" in 一句 and 备份[0] in 一句 for 一句 in 警告),
        "备份=%s 警告=%s" % (备份, 警告))
    备注 = ""
    if 备份:
        try:
            备份里 = __import__("pptx").Presentation(os.path.join(输出目录, "_备份", 备份[0]))
            备注 = list(备份里.slides)[1].notes_slide.notes_text_frame.text
        except Exception as 异常:
            备注 = "读不出来：" + type(异常).__name__
    检查("③b 备份里装的**正是老师改过的那版**（打开来读，不靠文件名猜）",
        "【探针】" in 备注, "备份备注=" + 备注[:24])
    新的稿 = [一页["备注"] for 一页 in 课件.读各页(路径)]
    检查("③c 新生成的那份**不含**老师的探针（备份出去之后确实换成了新的）",
        all("【探针】" not in 稿 for 稿 in 新的稿), str(新的稿)[:60])

    路径, 警告 = 生成一次("第三次b")
    检查("③d 备份过之后、没再改动 → **不再备份**（否则每生成一次多一份，还一直说你改过）",
        len(备份们()) == 1 and not any("不像是我们刚生成的那版" in 一句 for 一句 in 警告), str(备份们()))

    改备注(路径, "老师又改了一次【探针2】")
    路径, 警告 = 生成一次("第四次")
    检查("④ 同一分钟连改两次 → 两份备份都在（不会被彼此顶掉）", len(备份们()) == 2, str(备份们()))

    os.remove(_课件.生成记录路径(路径))
    路径, 警告 = 生成一次("第五次")
    检查("⑤ 没有生成记录 → 当\"不是邵新生成的\"，更要先备份并说明原因",
        len(备份们()) == 3 and any("不是邵新生成的" in 一句 for 一句 in 警告),
        "备份 %d 份；警告=%s" % (len(备份们()), 警告[:1]))

    顶层 = sorted(名 for 名 in os.listdir(输出目录)
                 if 名.endswith(".pptx") and not 名.startswith("_"))
    检查("⑥ 备份都在 _备份/ 里、顶层只剩正常课件（不会混进课件列表）",
        all("老师改过" not in 名 for 名 in 顶层) and bool(备份们()),
        "顶层=%s" % 顶层)

    原记录路径 = _课件.生成记录路径
    _课件.生成记录路径 = lambda p: 输出目录
    try:
        生成一次("第六次")
        检查("⑦ 生成记录写不出来 → 当场报人话", False, "居然默默过去了")
    except RuntimeError as 异常:
        检查("⑦ 生成记录写不出来 → 当场报人话（不许静默）",
            "生成记录写不出去" in str(异常), str(异常)[:56])
    finally:
        _课件.生成记录路径 = 原记录路径
finally:
    配置.取目录 = 原取目录
    shutil.rmtree(临时, ignore_errors=True)

for 序号, (名字, 通过, 说明) in enumerate(结果, 1):
    print("  %2d. %s %s —— %s" % (序号, "通过" if 通过 else "不过", 名字, 说明))
通过数 = sum(1 for _名, 通, _说 in 结果 if 通)
print("\n%d/%d" % (通过数, len(结果)))
sys.exit(0 if 通过数 == len(结果) else 1)
