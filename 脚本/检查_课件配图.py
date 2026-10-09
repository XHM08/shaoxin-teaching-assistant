
import json
import os
import shutil
import sys
import tempfile

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from 地基 import 大模型, 幻灯片, 课件, 配置
from 能力 import _课件

原取目录 = 配置.取目录
原读配置 = 大模型.读配置
原问 = 大模型.问
原读技能包 = __import__("地基.技能包", fromlist=["x"]).读技能包全文
临时 = tempfile.mkdtemp(prefix="课件配图自检_")
材料 = os.path.join(临时, "材料")
结果 = []


def 检查(名字, 通过, 说明=""):
    结果.append((名字, bool(通过), 说明))


def 画图(名, 大小=(120, 80), 在哪=None):
    from PIL import Image
    路径 = os.path.join(在哪 or 材料, 名)
    os.makedirs(os.path.dirname(路径), exist_ok=True)
    Image.new("RGB", 大小, (30, 90, 160)).save(路径)
    return 路径


提示词捕获 = []


def 假问(提示词, 服务商=None, **其余):
    提示词捕获.append({"提示词": 提示词, "图": 其余.get("图")})
    return json.dumps({
        "title": "配图自检",
        "pages": [
            {"kind": "封面", "title": "封面", "points": ["六年级"], "script": "开场", "图": ""},
            {"kind": "知识", "title": "看图的这一页", "关键句": "看这张图",
             "points": ["一条", "两条"], "script": "图里画着一块蓝色方块。", "图": "竖图.png"},
            {"kind": "练习", "title": "练一练", "points": ["算一算"], "script": "自己做"},
            {"kind": "小结", "title": "小结", "关键句": "记住", "points": ["一条"], "script": "收尾", "图": ""},
        ],
    }, ensure_ascii=False)


os.makedirs(材料)
画图("横图.png", (800, 400))
画图("竖图.png", (300, 900))
画图("子目录/嵌套图.jpg", (200, 200))
try:
    from PIL import Image
    Image.new("RGB", (100, 100), (10, 10, 10)).save(os.path.join(材料, "网页存下来的.webp"))
except Exception:
    pass
with open(os.path.join(材料, "说明.md"), "w", encoding="utf-8") as 文件:
    文件.write("这不是图")

配置.取目录 = lambda 名: 材料 if 名 == "材料" else 原取目录(名)
大模型.问 = 假问
__import__("地基.技能包", fromlist=["x"]).读技能包全文 = lambda 代号: "（自检用的假技能包）"

try:
    清单, 说明 = _课件.可配图()
    检查("① 可配图只列插得进的图（webp / md 不在内）",
       清单 == ["横图.png", "竖图.png", "子目录/嵌套图.jpg"],
       "实际：" + "、".join(清单))
    检查("② 清单是相对 材料/ 的路径、且用 / 分隔", all("/" in 名 for 名 in 清单 if "子" in 名),
       "实际：" + "、".join(清单))

    许多 = os.path.join(临时, "多图材料")
    os.makedirs(许多)
    for 号 in range(8):
        画图("图%d.png" % 号, (60, 60), 在哪=许多)
    老目录 = 材料
    配置.取目录 = lambda 名: 许多 if 名 == "材料" else 原取目录(名)
    多清单, 多说明 = _课件.可配图()
    配置.取目录 = lambda 名: 老目录 if 名 == "材料" else 原取目录(名)
    检查("③ 超过上限时只发前几张、而且说清楚了", len(多清单) == _课件.最多配图 and "只发前" in 多说明,
       "发了 %d 张；说明：%s" % (len(多清单), 多说明[:40]))
    检查("④ 上限没超过大模型自己的上限（一次最多几张图）",
       _课件.最多配图 <= 大模型.一次最多几张图,
       "最多配图=%d，大模型上限=%d" % (_课件.最多配图, 大模型.一次最多几张图))
    检查("⑤ 可插图后缀 ⊆ 大模型能发的后缀（两个集合不该相等）",
       set(幻灯片.可插图后缀) <= set(大模型.图片类型),
       "只插得进：%s；能发给模型：%s" % (tuple(幻灯片.可插图后缀), tuple(大模型.图片类型)))

    大模型.读配置 = lambda *a, **k: {"default": "看图的",
        "providers": {"看图的": {"base_url": "https://例子.invalid", "model": "看图的模型",
                              "key_env": "假的", "支持图片": True}}}
    上下文 = _课件.上下文表()
    数据 = _课件.生成("自检技能包", "分数的除法", "4", 上下文)
    捕获 = 提示词捕获[-1]
    图们 = 捕获["图"] or []
    检查("⑥ 标了支持图片 → 图真的随请求发出去", len(图们) == 3, "实际发了 %d 张" % len(图们))
    检查("⑦ 发出去的是（文件名标签, 路径）成对，标签就是清单里的名字",
       bool(图们) and all(isinstance(一条, tuple) and 一条[0].startswith("文件名：")
                         and 一条[0][4:] in 清单 for 一条 in 图们),
       "实际：" + str([一条[0] for 一条 in 图们 if isinstance(一条, tuple)])[:80])
    检查("⑧ 提示词里写了「已经附在提示词后面」", "已经附在提示词后面" in 捕获["提示词"])
    检查("⑨ 提示词里列了文件名，且要求照抄", "横图.png" in 捕获["提示词"] and "照抄" in 捕获["提示词"])
    检查("⑩ 提示词里说了只有 知识 / 例题 两页能配图",
       "只有【知识】【例题】两页可以填" in 捕获["提示词"])
    检查("⑪ 这次真发了图 → 警告里说的是「随提示词发了 N 张」",
       any("随提示词发了" in 一句 for 一句 in 上下文["警告"]) and
       not any("没标「支持图片」" in 一句 for 一句 in 上下文["警告"]),
       "警告：" + "；".join(上下文["警告"])[:100])

    大模型.读配置 = lambda *a, **k: {"default": "纯文本的",
        "providers": {"纯文本的": {"base_url": "https://例子.invalid", "model": "纯文本",
                                "key_env": "假的"}}}
    提示词捕获.clear()
    上下文2 = _课件.上下文表()
    _课件.生成("自检技能包", "分数的除法", "4", 上下文2)
    捕获2 = 提示词捕获[-1]
    检查("⑫ 没标支持图片 → 一张图也不发", not 捕获2["图"], "实际：" + str(捕获2["图"])[:60])
    检查("⑬ 没标支持图片 → 提示词里不列文件名（免得它照抄名字配上一张没看过的图）",
       "横图.png" not in 捕获2["提示词"] and "没有附上任何图片" in 捕获2["提示词"])
    检查("⑭ 没标支持图片 → 警告里告诉了老师（不静默降级）",
       any("没标「支持图片」" in 一句 for 一句 in 上下文2["警告"]),
       "警告：" + "；".join(上下文2["警告"])[:100])

    大模型.读配置 = lambda *a, **k: {"default": "看图的",
        "providers": {"看图的": {"base_url": "https://例子.invalid", "model": "看图的模型",
                              "key_env": "假的", "支持图片": True}}}
    输出目录 = os.path.join(临时, "课件")
    老输出 = 配置.取目录
    配置.取目录 = lambda 名: 输出目录 if 名 == "课件输出" else 老输出(名)
    上下文3 = _课件.上下文表()
    数据3 = _课件.生成("自检技能包", "分数的除法", "4", 上下文3)
    文件名 = _课件.保存pptx(数据3, "自检技能包", "分数的除法", 上下文3)
    路径3 = os.path.join(输出目录, 文件名)

    from pptx import Presentation
    prs = Presentation(路径3)
    页们 = list(prs.slides)
    图数 = sum(1 for 页 in 页们 for s in 页.shapes if hasattr(s, "image"))
    检查("⑮ 端到端：生成的 pptx 里真插进了图", 图数 == 1, "实际 %d 张图" % 图数)
    检查("⑯ 配上的页码回到调用方手里", 上下文3["配图页"] == [2],
       "实际：" + str(上下文3["配图页"]))
    检查("⑰ 讲稿仍在备注里（老约定没被配图打乱）",
       all(页.has_notes_slide and 页.notes_slide.notes_text_frame.text.strip() for 页 in 页们),
       "%d 页" % len(页们))
    检查("⑱ 图插在右半边（不压文字栏）",
       all((s.left + s.width) / 914400 >= 幻灯片.图左 - 0.02
           for 页 in 页们 for s in 页.shapes if hasattr(s, "image")))
    检查("⑲ 图的右缘贴住图区右边线（用竖图才验得出：横图居中与靠右结果一样）",
       all(abs((s.left + s.width) / 914400 - (幻灯片.图左 + 幻灯片.图宽)) < 0.03
           for 页 in 页们 for s in 页.shapes if hasattr(s, "image")))
    检查("⑳ 图不压页脚（下缘 ≤ 6.45）",
       all((s.top + s.height) / 914400 <= 6.45 + 0.02
           for 页 in 页们 for s in 页.shapes if hasattr(s, "image")))
    检查("㉑ 提示词里的文件清单与实际发出去的图一一对应（不多不少）",
       [一条[0][4:] for 一条 in 图们] == [名 for 名 in 清单] if 图们 else False)

    坏数据 = {"title": "编的", "pages": [
        {"kind": "知识", "title": "编的名字", "points": ["一条"], "script": "讲稿", "图": "根本没有这张.png"}]}
    上下文4 = _课件.上下文表()
    幻灯片.写出(坏数据, os.path.join(临时, "编的.pptx"), 上下文4)
    检查("㉒ 图名是编的 → 不崩、出页、且警告点名到页",
       os.path.isfile(os.path.join(临时, "编的.pptx")) and
       any("材料/ 里找不到这个文件" in 一句 for 一句 in 上下文4["警告"]) and
       上下文4["配图页"] == [],
       "警告：" + "；".join(上下文4["警告"])[:90])
finally:
    配置.取目录 = 原取目录
    大模型.读配置 = 原读配置
    大模型.问 = 原问
    __import__("地基.技能包", fromlist=["x"]).读技能包全文 = 原读技能包
    shutil.rmtree(临时, ignore_errors=True)

for 序号, (名字, 通过, 说明) in enumerate(结果, 1):
    print("  %2d. %s %s：%s" % (序号, "通过" if 通过 else "不过", 名字, 说明))
通过数 = sum(1 for _名, 通, _说 in 结果 if 通)
print("\n%d/%d" % (通过数, len(结果)))
sys.exit(0 if 通过数 == len(结果) else 1)
