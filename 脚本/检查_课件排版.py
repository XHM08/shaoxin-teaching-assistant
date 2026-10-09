
import os
import shutil
import sys
import tempfile

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from 地基 import 配置, 幻灯片

容差 = 0.06
画布 = 幻灯片.页宽
边距 = 幻灯片.边距

结果 = []


def 检查(名字, 通过, 说明=""):
    结果.append((名字, bool(通过), 说明))


夹具 = {
    "title": "分数的除法",
    "pages": [
        {"kind": "封面", "title": "分数的除法",
         "points": ["六年级 · 数学"], "script": "开场白"},
        {"kind": "知识", "title": "认识倒数",
         "关键句": "2/3 的倒数是 3/2",
         "points": ["分子分母交换位置", "0 没有倒数"], "script": "讲稿一"},
        {"kind": "例题", "title": "分数除以整数",
         "关键句": "3/4 ÷ 2 = 3/4 × 1/2 = 3/8",
         "points": ["先化除为乘", "再约分"], "script": "讲稿二"},
        {"kind": "练习", "title": "练一练",
         "关键句": "先独立完成，再看订正",
         "points": ["5/6 ÷ 1/3", "4/7 ÷ 2"], "script": "讲稿三"},
        {"kind": "小结", "title": "分数除法小结",
         "关键句": "除以一个数，等于乘它的倒数",
         "points": ["从整数推到分数", "注意 0 不能作除数"], "script": "讲稿四"},
    ],
}

怪样子 = {"title": "怪数据", "pages": [
    {"title": "整段字符串的要点", "关键句": "看看会不会一个字一行",
     "points": "第一条要点\n第二条要点\n第三条要点", "script": "讲稿"},
    {"title": "啥都没有的一页"},
    {"title": "kind 是没见过的类型", "kind": "随便写的", "points": ["应当按知识排"]},
]}

临时目录 = tempfile.mkdtemp(prefix="课件排版自检_")
try:
    from pptx import Presentation

    路径 = os.path.join(临时目录, "夹具.pptx")
    幻灯片.写出(夹具, 路径)
    prs = Presentation(路径)
    检查("① 五页都写出来了", len(prs.slides) == len(夹具["pages"]),
        "实际 %d 页" % len(prs.slides))

    页们 = list(prs.slides)

    封面右 = max((s.left + s.width) / 914400 for s in 页们[0].shapes)
    封面左 = min(s.left / 914400 for s in 页们[0].shapes)
    检查("② 封面是整幅的（左右都铺到边）",
        封面左 <= 容差 and 封面右 >= 画布 - 容差,
        "左 %.2f 右 %.2f（画布 %.2f）" % (封面左, 封面右, 画布))

    for 序号, 页 in enumerate(页们[1:], 2):
        右 = max((s.left + s.width) / 914400 for s in 页.shapes)
        左 = min(s.left / 914400 for s in 页.shapes)
        检查("③ 第 %d 页没有右边空白（右 %.2f ≥ %.2f）" % (序号, 右, 画布 - 边距 - 容差),
            右 >= 画布 - 边距 - 容差)
        检查("④ 第 %d 页没贴到左边缘（左 %.2f）" % (序号, 左), 左 >= 边距 - 容差)

    没有页脚 = []
    for 序号, 页 in enumerate(页们[1:], 2):
        if not any(abs(s.top / 914400 - 6.92) < 0.2 for s in 页.shapes if s.has_text_frame):
            没有页脚.append(序号)
    检查("⑤ 内容页都有页脚（第几页）", not 没有页脚,
        ("缺页脚的是第 " + "、".join(map(str, 没有页脚)) + " 页") if 没有页脚 else "")

    缺标题 = []
    for 序号, (页, 页数据) in enumerate(zip(页们, 夹具["pages"]), 1):
        标题 = str(页数据["title"])
        if not any(标题 in s.text_frame.text for s in 页.shapes if s.has_text_frame):
            缺标题.append(序号)
    检查("⑥ 每页都能找到标题", not 缺标题,
        ("缺的是第 " + "、".join(map(str, 缺标题)) + " 页") if 缺标题 else "")

    缺备注 = [序号 for 序号, 页 in enumerate(页们, 1)
             if not (页.has_notes_slide and 页.notes_slide.notes_text_frame.text.strip())]
    检查("⑦ 讲稿都进了备注（投屏不显示、演讲者视图能看到）", not 缺备注,
        ("缺的是第 " + "、".join(map(str, 缺备注)) + " 页") if 缺备注 else "")

    怪路径 = os.path.join(临时目录, "怪.pptx")
    try:
        幻灯片.写出(怪样子, 怪路径)
        怪prs = Presentation(怪路径)
        怪文本 = [s.text_frame.text for 页 in 怪prs.slides for s in 页.shapes
                 if s.has_text_frame]
        全部 = "\n".join(怪文本)
        检查("⑧ points 是一整段字符串时没有崩", True)
        检查("⑨ 那段字符串没被拆成一字一行",
            "第一条要点" in 全部 and "第二条要点" in 全部,
            "真实内容：" + 全部.replace("\n", " / ")[:120])
        检查("⑩ 少字段 / 没见过的 kind 也能出页", len(怪prs.slides) == 3,
            "实际 %d 页" % len(怪prs.slides))
        检查("⑪ 没见过的 kind 被当成知识页排（有页脚）",
            any(abs(s.top / 914400 - 6.92) < 0.2
                for s in 怪prs.slides[2].shapes if s.has_text_frame))
    except Exception as 异常:
        检查("⑧ points 是一整段字符串时没有崩", False,
            type(异常).__name__ + "：" + str(异常).splitlines()[0][:90])

    try:
        空路径 = os.path.join(临时目录, "空.pptx")
        幻灯片.写出({"title": "空", "pages": []}, 空路径)
        检查("⑫ 一个页面都没有时也产出可打开的文件（而不是 0 页）",
            len(Presentation(空路径).slides) >= 1)
    except Exception as 异常:
        检查("⑫ 一个页面都没有时也产出可打开的文件", False,
            type(异常).__name__ + "：" + str(异常).splitlines()[0][:90])
    try:
        from PIL import Image
        图目录 = os.path.join(临时目录, "材料")
        os.makedirs(图目录, exist_ok=True)
        Image.new("RGB", (800, 400), (30, 90, 160)).save(os.path.join(图目录, "横图.png"))
        Image.new("RGB", (300, 900), (160, 60, 60)).save(os.path.join(图目录, "竖图.png"))

        带图 = {"title": "配图", "pages": [
            {"kind": "封面", "title": "封面", "points": ["六年级"], "script": "开场", "图": "横图.png"},
            {"kind": "知识", "title": "横图", "关键句": "看这张图",
             "points": ["一条", "两条"], "script": "讲稿", "图": "横图.png"},
            {"kind": "例题", "title": "竖图", "关键句": "竖着的图",
             "points": ["一条"], "script": "讲稿", "图": "竖图.png"},
            {"kind": "练习", "title": "练习页填了图", "points": ["算一算"], "script": "讲稿", "图": "横图.png"},
            {"kind": "小结", "title": "小结", "关键句": "记住", "points": ["一条"], "script": "讲稿", "图": ""},
        ]}
        上下文 = {"图目录": 图目录, "警告": [], "配图页": []}
        图路径 = os.path.join(临时目录, "配图.pptx")
        幻灯片.写出(带图, 图路径, 上下文)
        图prs = Presentation(图路径)
        图页们 = list(图prs.slides)
        图片们 = [(序号, s) for 序号, 页 in enumerate(图页们, 1)
                 for s in 页.shapes if hasattr(s, "image")]

        检查("⑬ 只有能配图的页（知识/例题）出了图，封面/练习/小结都没有",
            [序号 for 序号, _ in 图片们] == [2, 3],
            "出图的页：" + str([序号 for 序号, _ in 图片们]))
        检查("⑭ 非配图页填了图会被忽略，而且出声告诉老师",
            any("第 1 页（封面）" in 一句 for 一句 in 上下文["警告"]) and
            any("第 4 页（练习）" in 一句 for 一句 in 上下文["警告"]),
            "警告：" + "；".join(上下文["警告"])[:100])
        检查("⑮ 图靠右：右缘贴住图区右边线（竖图也不留一道空）",
            all(abs((s.left + s.width) / 914400 - (幻灯片.图左 + 幻灯片.图宽)) < 0.03
                for _, s in 图片们),
            "右缘：" + "、".join("%.2f" % ((s.left + s.width) / 914400) for _, s in 图片们))
        检查("⑯ 图不压页脚（下缘 ≤ 6.45）",
            all((s.top + s.height) / 914400 <= 6.45 + 0.02 for _, s in 图片们),
            "下缘：" + "、".join("%.2f" % ((s.top + s.height) / 914400) for _, s in 图片们))
        比例 = [(s.width / 914400) / (s.height / 914400) for _, s in 图片们]
        检查("⑰ 长宽比没被拉变形（横图 2.0、竖图 0.33）",
            len(比例) == 2 and abs(比例[0] - 2.0) < 0.05 and abs(比例[1] - 300 / 900) < 0.05,
            "实际：" + "、".join("%.2f" % 比 for 比 in 比例))
        缺页脚的 = [序号 for 序号, 页 in enumerate(图页们[1:], 2)
                  if not any(abs(s.top / 914400 - 6.92) < 0.2
                             for s in 页.shapes if s.has_text_frame)]
        检查("⑱ 带图页的页脚还在（配图没把它挤掉）", not 缺页脚的,
            ("缺页脚的是第 " + "、".join(map(str, 缺页脚的)) + " 页") if 缺页脚的 else "")
    except ImportError as 异常:
        检查("⑬~⑱ 配图排版（要 Pillow 现场画图）", False,
            "ImportError：" + str(异常)[:60] + "（先看是不是这台电脑没有 Pillow）")
    except Exception as 异常:
        检查("⑬~⑱ 配图排版", False, type(异常).__name__ + "：" + str(异常).splitlines()[0][:90])

    try:
        提示词文本 = open(os.path.join(配置.取目录("提示词"), "课件生成.md"), encoding="utf-8").read()
        检查("⑲ 提示词里说的能配图的两页 == 代码里的 可配图的类型",
            幻灯片.可配图的类型 == ("知识", "例题")
            and "只有【知识】【例题】两页可以填" in 提示词文本,
            "代码：" + str(幻灯片.可配图的类型))
    except Exception as 异常:
        检查("⑲ 提示词与代码的两页约定一致", False, type(异常).__name__ + "：" + str(异常)[:80])

finally:
    shutil.rmtree(临时目录, ignore_errors=True)

坏的 = 0
print()
for 名字, 通过, 说明 in 结果:
    print(("  通过  " if 通过 else "  失败  ") + 名字 + ("　→　" + 说明 if 说明 else ""))
    坏的 += 0 if 通过 else 1
print("全部 %d 项通过" % len(结果) if not 坏的 else "%d 项失败" % 坏的)
sys.exit(0 if not 坏的 else 1)
