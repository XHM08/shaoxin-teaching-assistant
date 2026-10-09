
import http.server
import io
import json
import os
import shutil
import sys
import tempfile
import threading
import wave

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from 地基 import 幻灯片, 审计, 课件, 配置, 语音
from 能力 import 讲稿配音

原取目录 = 配置.取目录
原取地址 = 语音.取地址
原记一笔 = 审计.记一笔
临时 = tempfile.mkdtemp(prefix="讲稿配音自检_")
课件目录 = os.path.join(临时, "课件输出")
参考音目录 = os.path.join(临时, "参考音")
结果 = []
记录 = []
每秒每字 = 0.05


def 检查(名字, 通过, 说明=""):
    结果.append((名字, bool(通过), 说明))


class 假语音(http.server.BaseHTTPRequestHandler):
    def do_POST(self):
        长 = int(self.headers.get("Content-Length") or 0)
        体 = json.loads(self.rfile.read(长).decode("utf-8"))
        记录.append({"路径": self.path, "体": 体})
        率 = 8000
        帧 = max(1, int(max(0.05, len(体.get("text", "")) * 每秒每字) * 率))
        缓冲 = io.BytesIO()
        with wave.open(缓冲, "wb") as 音频:
            音频.setnchannels(1)
            音频.setsampwidth(2)
            音频.setframerate(率)
            音频.writeframes(b"\x00\x00" * 帧)
        数据 = 缓冲.getvalue()
        self.send_response(200)
        self.send_header("Content-Type", "audio/wav")
        self.send_header("Content-Length", str(len(数据)))
        self.end_headers()
        self.wfile.write(数据)

    def log_message(self, *其余):
        pass


os.makedirs(课件目录)
os.makedirs(参考音目录)

with wave.open(os.path.join(参考音目录, "参考音.wav"), "wb") as 音频:
    音频.setnchannels(1)
    音频.setsampwidth(2)
    音频.setframerate(16000)
    音频.writeframes(b"\x00\x00" * 16000)
with open(os.path.join(参考音目录, "参考音.txt"), "w", encoding="utf-8") as 文件:
    文件.write("这是自检用的参考音，念的是这一句话\n")

夹具 = {"title": "配音夹具", "pages": [
    {"kind": "封面", "title": "封面", "points": ["六年级"], "script": "同学们好，今天讲分数的除法"},
    {"kind": "知识", "title": "这页没有讲稿", "关键句": "x", "points": ["一条"]},
    {"kind": "小结", "title": "小结", "关键句": "记住", "points": ["一条"], "script": "下课之前再记一遍结论"},
]}
幻灯片.写出(夹具, os.path.join(课件目录, "配音夹具.pptx"))

配置.取目录 = lambda 名: {"课件输出": 课件目录, "参考音": 参考音目录}.get(名) or 原取目录(名)

服务 = http.server.ThreadingHTTPServer(("127.0.0.1", 0), 假语音)
端口 = 服务.server_address[1]
线程 = threading.Thread(target=服务.serve_forever, daemon=True)
线程.start()
语音.取地址 = lambda: "http://127.0.0.1:%d" % 端口

收到 = []
审计.记一笔 = lambda 类别, 内容, **其余: 收到.append((类别, dict(内容)))

try:
    try:
        语音.取地址 = lambda: "http://127.0.0.1:1"
        讲稿配音.处理("配音夹具.pptx")
        检查("① 服务没起 → 给人话", False, "居然没报错")
    except Exception as 异常:
        检查("① 服务没起 → 给人话",
            "语音服务没在跑" in str(异常) and "启动" in str(异常),
            type(异常).__name__ + "：" + str(异常).splitlines()[0][:70])
    语音.取地址 = lambda: "http://127.0.0.1:%d" % 端口

    老参考 = 配置.取目录
    空目录 = os.path.join(临时, "空参考音")
    os.makedirs(空目录)
    配置.取目录 = lambda 名: 空目录 if 名 == "参考音" else 老参考(名)
    try:
        讲稿配音.处理("配音夹具.pptx")
        检查("② 没放参考音 → 给人话", False, "居然没报错")
    except ValueError as 异常:
        检查("② 没放参考音 → 给人话（且带上合规提醒）",
            "参考音" in str(异常) and "本人同意" in str(异常),
            str(异常).replace("\n", " ")[:70])
    配置.取目录 = 老参考

    记录.clear()
    回执 = 讲稿配音.处理("配音夹具.pptx", "1.0")
    配音目录 = os.path.join(课件目录, "配音夹具_配音")
    文件们 = sorted(os.listdir(配音目录)) if os.path.isdir(配音目录) else []
    波形 = [名 for 名 in 文件们 if 名.endswith(".wav")]
    检查("③ 有讲稿的页各出一个 wav（第 2 页没讲稿 → 跳过）",
        波形 == ["第1页.wav", "第3页.wav"], "实际：" + "、".join(文件们))
    检查("③b 配音目录里有台账（差分就靠它，改一页只重配一页）",
        讲稿配音.配音台账名 in 文件们, "实际：" + "、".join(文件们))
    检查("④ 没讲稿的页在回执里说明白了（不是悄悄少一页）",
        "1 页没有讲稿" in 回执["提示"] and "第 2 页" in 回执["提示"], 回执["提示"][:90])

    检查("⑤ 合成用的文本就是备注里的讲稿，且按页序",
        [记["体"]["text"] for 记 in 记录] ==
        ["同学们好，今天讲分数的除法", "下课之前再记一遍结论"],
        "实际：" + str([记["体"]["text"][:12] for 记 in 记录]))

    体 = 记录[0]["体"] if 记录 else {}
    检查("⑥ 请求体字段齐全且对（ref_audio_path / prompt_text / lang / speed / wav）",
        os.path.basename(str(体.get("ref_audio_path"))) == "参考音.wav"
        and 体.get("prompt_text", "").startswith("这是自检用的参考音")
        and 体.get("text_lang") == "zh" and 体.get("prompt_lang") == "zh"
        and float(体.get("speed_factor", 0)) == 1.0 and 体.get("media_type") == "wav"
        and all(记["路径"] == "/tts" for 记 in 记录),
        "实际：" + json.dumps(体, ensure_ascii=False)[:110])

    def 秒数(名):
        with wave.open(os.path.join(配音目录, 名), "rb") as 音频:
            return 音频.getnframes() / float(音频.getframerate())
    期望1 = len("同学们好，今天讲分数的除法") * 每秒每字
    期望3 = len("下课之前再记一遍结论") * 每秒每字
    检查("⑦ 每页音频的时长与那页讲稿的字数对得上（页与稿没串）",
        abs(秒数("第1页.wav") - 期望1) < 0.15 and abs(秒数("第3页.wav") - 期望3) < 0.15,
        "第1页 %.2fs(期望 %.2f) / 第3页 %.2fs(期望 %.2f)"
        % (秒数("第1页.wav"), 期望1, 秒数("第3页.wav"), 期望3))

    检查("⑧ 回执条目带页号、字数、时长、文件名",
        [(条["页"], 条["字数"]) for 条 in 回执["条目"]] == [(1, 13), (3, 10)],
        json.dumps(回执["条目"], ensure_ascii=False)[:110])

    检查("⑨ 审计记了「讲稿配音」一笔（含页数与字数）",
        any(类别 == "讲稿配音" and 内容.get("合成页数") == 2 for 类别, 内容 in 收到),
        json.dumps(收到[-1][1], ensure_ascii=False)[:90] if 收到 else "没记")

    try:
        讲稿配音.处理("../../配置.json")
        检查("⑩ 课件文件名带目录 → 拦住", False, "居然放过去了")
    except ValueError as 异常:
        检查("⑩ 课件文件名带目录 → 拦住并给人话", "只填文件名" in str(异常), str(异常)[:60])

    try:
        讲稿配音.处理("没这份.pptx")
        检查("⑪ 课件不存在 → 给人话并列出已有的", False, "居然没报错")
    except ValueError as 异常:
        检查("⑪ 课件不存在 → 给人话并列出已有的",
            "没有这份课件" in str(异常) and "配音夹具.pptx" in str(异常),
            str(异常).replace("\n", " ")[:80])

    try:
        讲稿配音.处理("配音夹具.pptx", "9")
        检查("⑫ 语速越界 → 拦住", False, "居然放过去了")
    except ValueError as 异常:
        检查("⑫ 语速越界 → 拦住并给人话", "语速只在" in str(异常), str(异常)[:60])

    目录名, 配音条目 = 讲稿配音.合成整份(os.path.join(课件目录, "配音夹具.pptx"))
    检查("⑬ 合成整份 返回的音频地址形状对（/配音/<目录>/第N页.wav，未编码）",
        [条["音频"] for 条 in 配音条目] == ["/配音/配音夹具_配音/第1页.wav",
                                        "/配音/配音夹具_配音/第3页.wav"],
        "实际：" + str([条["音频"] for 条 in 配音条目]))

    警告 = []
    讲稿配音.合成整份(os.path.join(课件目录, "配音夹具.pptx"), 1.0, 警告)
    检查("⑭ 没讲稿的页要在警告里说出来（调用方才有话告诉老师）",
        any("没有讲稿" in 一句 for 一句 in 警告), "警告：" + str(警告)[:80])

    语音.取地址 = lambda: "http://127.0.0.1:1"
    行不行, 原因 = 讲稿配音.能配音吗()
    检查("⑮ 服务没起 → 不能配音，且原因是「服务没在跑」+ 怎么启动",
        行不行 is False and "语音服务没在跑" in 原因 and "启动" in 原因, "原因：" + 原因[:70])
    语音.取地址 = lambda: "http://127.0.0.1:%d" % 端口
    老取 = 配置.取目录
    配置.取目录 = lambda 名: 空目录 if 名 == "参考音" else 老取(名)
    行不行, 原因 = 讲稿配音.能配音吗()
    配置.取目录 = 老取
    检查("⑯ 参考音没放 → 不能配音，且原因指到参考音",
        行不行 is False and "参考音" in 原因, "原因：" + 原因[:70])

    副本名 = "自检用-差分副本.pptx"
    副本路径 = os.path.join(课件目录, 副本名)
    shutil.copyfile(os.path.join(课件目录, "配音夹具.pptx"), 副本路径)
    副本配音目录 = os.path.join(课件目录, "自检用-差分副本_配音")

    def 配一次(语速=1.0):
        记录.clear()
        警告 = []
        讲稿配音.合成整份(副本路径, 语速, 警告)
        return len(记录), "；".join(警告)

    try:
        次, 话 = 配一次()
        检查("⑱ 第一次：有讲稿的页各调一次合成（2 页 → 2 次）", 次 == 2, "调了 %d 次；%s" % (次, 话))

        次, 话 = 配一次()
        检查("⑲ 没改过再配一次：一次都不调（沿用上次），且话说明了",
            次 == 0 and "一页都没重配" in 话, "调了 %d 次；%s" % (次, 话))

        from pptx import Presentation
        演示 = Presentation(副本路径)
        list(演示.slides)[2].notes_slide.notes_text_frame.text = "换一段讲稿试试，这页已经改过了"
        演示.save(副本路径)
        次, 话 = 配一次()
        检查("⑳ 只改一页备注 → 只重配那一页（1 次），并点名第 3 页、说清其余沿用",
            次 == 1 and "第 3 页" in 话 and "其余 1 页" in 话, "调了 %d 次；%s" % (次, 话))

        演示 = Presentation(副本路径)
        for 形 in list(演示.slides)[0].shapes:
            if 形.has_text_frame and 形.text_frame.text.strip():
                形.text_frame.text = "（老师改的写法）先拿整数试一试"
                break
        演示.save(副本路径)
        次, 话 = 配一次()
        检查("㉑ 文字改了备注没改 → 不重配，但点名提醒「讲稿（备注）没动」",
            次 == 0 and "第 1 页" in 话 and "讲稿（备注）没动" in 话, "调了 %d 次；%s" % (次, 话))

        次, 话 = 配一次(1.2)
        检查("㉒ 语速变了 → 全部重配（换个速度是另一条音频，不能拿旧的凑）", 次 == 2,
            "调了 %d 次；%s" % (次, 话))

        with open(os.path.join(副本配音目录, 讲稿配音.配音台账名), "w", encoding="utf-8") as 文件:
            文件.write("{这不是 json")
        次, 话 = 配一次(1.2)
        检查("㉓ 台账坏掉 → 不崩，当成没有台账全量重配", 次 == 2, "调了 %d 次；%s" % (次, 话))

        检查("㉔ 配音只读不写：老师的改动还在课件里",
            课件.读各页(副本路径)[2]["备注"].startswith("换一段讲稿试试"),
            课件.读各页(副本路径)[2]["备注"][:20])
    finally:
        shutil.rmtree(副本配音目录, ignore_errors=True)
        if os.path.isfile(副本路径):
            os.remove(副本路径)

    检查("⑰ 要不要配音：填「不」才跳过，空着/「要」都要配",
        (讲稿配音.要配音吗("不") is False and 讲稿配音.要配音吗("不要") is False
         and 讲稿配音.要配音吗("") is True and 讲稿配音.要配音吗("要") is True
         and 讲稿配音.要配音吗(None) is True),
        "「不」→%s 空→%s「要」→%s" % (讲稿配音.要配音吗("不"), 讲稿配音.要配音吗(""),
                                 讲稿配音.要配音吗("要")))
finally:
    服务.shutdown()
    配置.取目录 = 原取目录
    语音.取地址 = 原取地址
    审计.记一笔 = 原记一笔
    shutil.rmtree(临时, ignore_errors=True)

for 序号, (名字, 通过, 说明) in enumerate(结果, 1):
    print("  %2d. %s %s：%s" % (序号, "通过" if 通过 else "不过", 名字, 说明))
通过数 = sum(1 for _名, 通, _说 in 结果 if 通)
print("\n%d/%d" % (通过数, len(结果)))
sys.exit(0 if 通过数 == len(结果) else 1)
